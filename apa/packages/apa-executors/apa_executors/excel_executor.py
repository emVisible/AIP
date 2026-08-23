"""apa-executors: Excel 操作执行器（对标影刀 Excel 指令组）。

动作域：excel.*
依赖：openpyxl（optional-dependency: apa-core[excel]）

支持：
  excel.read_range     读取区域 → 表格数据
  excel.append_row     追加一行
  excel.write_cell     写单元格
  excel.get_formula    读公式文本
  excel.set_formula    写公式
  excel.list_sheets    列出工作表
  excel.add_sheet      新建工作表

安全：所有路径限定在进程可达文件系统；写操作先落盘临时文件再原子替换。
"""
from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

from apa_sdk.executor_base import AIPExecutor


def _load_openpyxl():
    try:
        import openpyxl  # noqa: PLC0415
        return openpyxl
    except ImportError as e:
        raise RuntimeError(
            "openpyxl is required for excel.* actions: "
            "uv pip install openpyxl (or pip install 'apa-core[excel]')",
        ) from e


class ExcelExecutor(AIPExecutor):
    """Excel 工作簿读写执行器。"""

    def start_observation(self) -> None:
        pass

    # ---- 内部工具 ----------------------------------------------------------

    @staticmethod
    def _wb(path: str, read_only: bool = False, data_only: bool = False):
        openpyxl = _load_openpyxl()
        p = Path(path)
        if not p.is_file():
            raise FileNotFoundError(f"workbook not found: {path}")
        return openpyxl.load_workbook(
            str(p), read_only=read_only, data_only=data_only)

    @staticmethod
    def _atomic_save(wb, path: str) -> None:
        """临时目录写出 + 原子替换，避免写坏原文件。"""
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(path) or ".")
        os.close(fd)
        try:
            wb.save(tmp)
            os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):  # save 失败时清理
                os.unlink(tmp)

    def _execute_action(self, name: str, params: dict) -> Tuple[bool, dict]:
        handlers = {
            "excel.read_range": self._read_range,
            "excel.append_row": self._append_row,
            "excel.write_cell": self._write_cell,
            "excel.get_formula": self._get_formula,
            "excel.set_formula": self._set_formula,
            "excel.list_sheets": self._list_sheets,
            "excel.add_sheet": self._add_sheet,
        }
        fn = handlers.get(name)
        if fn is None:
            return False, {"code": "action_not_supported_by_executor"}
        try:
            return fn(params)
        except FileNotFoundError as e:
            return False, {"code": "file_not_found", "detail": str(e)}
        except RuntimeError as e:  # openpyxl 缺失
            return False, {"code": "dependency_missing", "detail": str(e)}
        except KeyError as e:
            return False, {"code": "missing_param", "detail": f"{e}"}
        except Exception as e:
            return False, {"code": type(e).__name__, "detail": str(e)}

    # ---- 动作实现 ----------------------------------------------------------

    def _read_range(self, p: dict) -> Tuple[bool, dict]:
        """读取 sheet!A1:D10 或整个 sheet 的已用区域。"""
        wb = self._wb(p["file"], read_only=True, data_only=True)
        try:
            ws = wb[p["sheet"]] if p.get("sheet") else wb.active
            rng = p.get("range")
            if rng:
                # 只读模式不支持切片，转换为行列边界
                from openpyxl.utils import range_boundaries
                min_col, min_row, max_col, max_row = range_boundaries(rng)
                rows_iter = ws.iter_rows(
                    min_row=min_row, max_row=max_row,
                    min_col=min_col, max_col=max_col,
                    values_only=True,
                )
            else:
                rows_iter = ws.iter_rows(values_only=True)
            header = p.get("header", True)
            rows = [list(r) for r in rows_iter]
            columns = [str(c) for c in rows[0]] if (header and rows) else []
            body = rows[1:] if (header and rows) else rows
            return True, {
                "columns": columns,
                "rows": body,
                "count": len(body),
                "sheet": ws.title,
            }
        finally:
            wb.close()

    def _append_row(self, p: dict) -> Tuple[bool, dict]:
        values = p["values"]
        wb = self._wb(p["file"])
        try:
            ws = wb[p["sheet"]] if p.get("sheet") else wb.active
            ws.append(values)
            self._atomic_save(wb, p["file"])
            return True, {
                "appended": len(values),
                "row_index": ws.max_row,
                "sheet": ws.title,
            }
        finally:
            wb.close()

    def _write_cell(self, p: dict) -> Tuple[bool, dict]:
        wb = self._wb(p["file"])
        try:
            ws = wb[p["sheet"]] if p.get("sheet") else wb.active
            ws[p["cell"]] = p["value"]
            self._atomic_save(wb, p["file"])
            return True, {"cell": p["cell"], "value": p["value"]}
        finally:
            wb.close()

    def _get_formula(self, p: dict) -> Tuple[bool, dict]:
        wb = self._wb(p["file"], data_only=False)
        try:
            ws = wb[p["sheet"]] if p.get("sheet") else wb.active
            return True, {"cell": p["cell"], "formula": ws[p["cell"]].value}
        finally:
            wb.close()

    def _set_formula(self, p: dict) -> Tuple[bool, dict]:
        formula: str = p["formula"]
        if not formula.startswith("="):
            formula = f"={formula}"
        wb = self._wb(p["file"])
        try:
            ws = wb[p["sheet"]] if p.get("sheet") else wb.active
            ws[p["cell"]] = formula
            self._atomic_save(wb, p["file"])
            return True, {"cell": p["cell"], "formula": formula}
        finally:
            wb.close()

    def _list_sheets(self, p: dict) -> Tuple[bool, dict]:
        wb = self._wb(p["file"], read_only=True)
        try:
            names = list(wb.sheetnames)
            return True, {"sheets": names, "active": wb.active.title}
        finally:
            wb.close()

    def _add_sheet(self, p: dict) -> Tuple[bool, dict]:
        wb = self._wb(p["file"])
        try:
            if p["sheet"] in wb.sheetnames:
                return False, {"code": "sheet_exists", "sheet": p["sheet"]}
            idx: Optional[int] = p.get("index")
            ws = wb.create_sheet(p["sheet"], index=idx)
            self._atomic_save(wb, p["file"])
            return True, {"sheet": ws.title, "position": (
                wb.index(ws) if idx is not None else None)}
        finally:
            wb.close()
